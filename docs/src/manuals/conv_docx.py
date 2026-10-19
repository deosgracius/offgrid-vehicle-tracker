import re, os
from io import BytesIO
from urllib.parse import unquote
from bs4 import BeautifulSoup
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from PIL import Image

BASE = os.path.dirname(os.path.abspath(__file__))
OUT = r"C:\Users\deogr\Documents\Fall Project\Off-Grid Vehicle dection project\docs\reports\Group9_Technical_Report.docx"

soup = BeautifulSoup(open(os.path.join(BASE, "REPORT.html"), encoding="utf-8").read(), "html.parser")
doc = Document()
for sec in doc.sections:
    sec.top_margin = sec.bottom_margin = Inches(0.9)
    sec.left_margin = sec.right_margin = Inches(1.0)
normal = doc.styles["Normal"]
normal.font.name = "Times New Roman"
normal.font.size = Pt(12)

def img_path(src):
    if src.startswith("file:///"):
        return unquote(src[8:]).replace("/", os.sep)
    return os.path.join(BASE, src)

def _clean(path):
    im = Image.open(path).convert("RGB")
    buf = BytesIO(); im.save(buf, format="PNG"); buf.seek(0)
    return buf, im.size

def add_runs(par, node):
    for child in node.children:
        if isinstance(child, str):
            if child.strip("\n") != "" or child == " ":
                par.add_run(re.sub(r"\s+", " ", child))
        elif child.name in ("b", "strong"):
            r = par.add_run(child.get_text()); r.bold = True
        elif child.name in ("i", "em"):
            r = par.add_run(child.get_text()); r.italic = True
        elif child.name == "code":
            r = par.add_run(child.get_text()); r.font.name = "Courier New"; r.font.size = Pt(10.5)
        elif child.name == "sup":
            r = par.add_run(child.get_text()); r.font.superscript = True
        elif child.name == "br":
            par.add_run().add_break()
        else:
            par.add_run(re.sub(r"\s+", " ", child.get_text()))

def body_par(node):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    add_runs(p, node)

def heading(text, size, before=14, center=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(4)
    if center: p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text); r.bold = True; r.font.size = Pt(size)

def centered_image(path, max_w=6.3, eq=False):
    if not os.path.exists(path):
        print("MISSING IMAGE:", path); return
    buf, (w, h) = _clean(path)
    aspect = w / h
    if eq:
        height = 0.42; width = height * aspect
        if width > max_w: width = max_w
    else:
        if aspect > 1.6:   width = min(max_w, 6.3)
        elif aspect > 0.8: width = 4.6
        else:              width = 2.4
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(buf, width=Inches(width))

def caption(node):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(re.sub(r"\s+", " ", node.get_text()).strip())
    r.font.size = Pt(10.5)

def table_from(node):
    rows = node.find_all("tr")
    ncols = max(len(r.find_all(["td", "th"])) for r in rows)
    t = doc.add_table(rows=0, cols=ncols)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r in rows:
        cells = r.find_all(["td", "th"])
        row = t.add_row()
        for i, c in enumerate(cells):
            cp = row.cells[i].paragraphs[0]
            run = cp.add_run(re.sub(r"\s+", " ", c.get_text()).strip())
            run.font.size = Pt(10); run.font.name = "Times New Roman"
            if c.name == "th": run.bold = True
    doc.add_paragraph().paragraph_format.space_after = Pt(0)

# ---- cover ----
cover = soup.find("div", class_="cover")
inst = cover.find("div", class_="inst")
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
lines = inst.get_text("\n").split("\n")
for i, line in enumerate(lines):
    r = p.add_run(line.strip()); r.font.size = Pt(12.5)
    if i < len(lines) - 1: r.add_break()
doc.add_paragraph()
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run(re.sub(r"\s+", " ", cover.find("h1").get_text()).strip())
r.bold = True; r.font.size = Pt(21)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run(cover.find("div", class_="sub").get_text().strip()); r.italic = True; r.font.size = Pt(13)
for _ in range(6): doc.add_paragraph()
auth = cover.find("div", class_="auth")
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
lines = [l.strip() for l in auth.get_text("\n").split("\n") if l.strip()]
for i, line in enumerate(lines):
    r = p.add_run(line); r.font.size = Pt(12)
    if i < len(lines) - 1: r.add_break()
doc.add_page_break()

# ---- body ----
page = soup.find("div", class_="page")
for el in page.children:
    if el.name is None:
        continue
    if el.name == "h2":
        txt = re.sub(r"\s+", " ", el.get_text()).strip()
        if txt.startswith("References") or txt.startswith("Appendix"):
            doc.add_page_break()
        heading(txt, 15, before=16)
    elif el.name == "p":
        body_par(el)
    elif el.name == "img":
        centered_image(img_path(el.get("src")), eq="eqimg" in (el.get("class") or []))
    elif el.name == "table":
        table_from(el)
    elif el.name == "ol":
        is_refs = "refs" in (el.get("class") or [])
        for n, li in enumerate(el.find_all("li"), 1):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.4)
            p.paragraph_format.first_line_indent = Inches(-0.25)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if not is_refs else WD_ALIGN_PARAGRAPH.LEFT
            lead = p.add_run(f"[{n}] " if is_refs else f"{n}. ")
            lead.bold = not is_refs
            add_runs(p, li)
    elif el.name == "div":
        cls = el.get("class") or []
        if "abstract" in cls:
            for sub in el.find_all("p"):
                body_par(sub)
        elif "fig" in cls:
            imgs = el.find_all("img")
            if len(imgs) == 2:
                p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for im in imgs:
                    pth = img_path(im.get("src"))
                    if os.path.exists(pth):
                        buf, _ = _clean(pth)
                        p.add_run().add_picture(buf, width=Inches(2.35))
                        p.add_run("  ")
            else:
                for im in imgs:
                    centered_image(img_path(im.get("src")))
            cap = el.find("div", class_="cap")
            if cap: caption(cap)
        elif "cap" in cls:
            caption(el)
        elif not cls:
            # dedicated figure page
            doc.add_page_break()
            for sub in el.children:
                if sub.name == "h2":
                    heading(re.sub(r"\s+", " ", sub.get_text()).strip(), 15, center=True)
                elif sub.name == "div" and "fig" in (sub.get("class") or []):
                    for im2 in sub.find_all("img"):
                        centered_image(img_path(im2.get("src")))
                    cap2 = sub.find("div", class_="cap")
                    if cap2: caption(cap2)
                elif sub.name == "p":
                    body_par(sub)
            doc.add_page_break()

doc.save(OUT)
d2 = Document(OUT)
print("saved:", OUT)
print("paragraphs:", len(d2.paragraphs), "tables:", len(d2.tables), "images:", len(d2.inline_shapes))
